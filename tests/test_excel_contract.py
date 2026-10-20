"""Plain-python tests for the excel-to-data-contract skill: python tests/test_excel_contract.py"""
import io, shutil, sys, tempfile, contextlib
from pathlib import Path
import openpyxl, yaml
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "skills/excel-to-data-contract/scripts"))
sys.path.insert(0, str(ROOT / "skills/contract-driven-pipeline/scripts"))
import excel_contract as xc  # noqa: E402
import contract_tools as ct  # noqa: E402

V, S = ROOT / "demo/contract_versions", ROOT / "demo/spec_versions"
tmp = Path(tempfile.mkdtemp())
load = lambda p: yaml.safe_load(Path(p).read_text())
quiet = lambda *a: contextlib.redirect_stdout(io.StringIO())

def run(args):
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        rc = xc.main(args)
    return rc, buf.getvalue()

# 1. round trip: Excel spec in specs/ == contract v1.0.0
rc, out = run(["convert", str(ROOT / "specs/cargo_shipments.xlsx"), str(tmp / "new.yaml")])
assert rc == 0 and "WRITTEN" in out, out
assert ct._canon(load(tmp / "new.yaml")) == ct._canon(load(V / "cargo_shipments.v1.0.0.yaml"))
assert not [f for f in ct.validate_contract(load(tmp / "new.yaml")) if f.level == "ERROR"]
# 2. same Excel again -> unchanged
rc, out = run(["convert", str(ROOT / "specs/cargo_shipments.xlsx"), str(tmp / "new.yaml")])
assert rc == 0 and "UNCHANGED" in out, out
# 3. additive change with blank version -> 1.1.0, identical to the demo v1.1.0 contract
rc, out = run(["convert", str(S / "cargo_shipments.v1.1.0.xlsx"), str(tmp / "new.yaml")])
assert rc == 0 and "proposed 1.1.0" in out, out
assert ct.diff_contracts(load(V / "cargo_shipments.v1.1.0.yaml"), load(tmp / "new.yaml"))["overall"] == "NONE"
assert not [f for f in ct.check_pipeline(load(tmp / "new.yaml"), ROOT / "expected_output/v1.1.0/transformations/cargo_shipments.py") if f.level == "ERROR"]
# 4. breaking change -> blocked, file untouched; with sign-off -> 2.0.0
before = (tmp / "new.yaml").read_text()
rc, out = run(["convert", str(S / "cargo_shipments.v2.0.0.xlsx"), str(tmp / "new.yaml")])
assert rc == 1 and "BREAKING" in out and (tmp / "new.yaml").read_text() == before, out
rc, out = run(["convert", str(S / "cargo_shipments.v1.2.0-mislabelled.xlsx"), str(tmp / "new.yaml")])
assert rc == 1 and "needs major" in out, out
rc, out = run(["convert", str(S / "cargo_shipments.v2.0.0.xlsx"), str(tmp / "new.yaml"), "--allow-breaking"])
assert rc == 0 and load(tmp / "new.yaml")["version"] == "2.0.0", out
# 5. messy Excel -> precise cell-level errors, nothing written
bad = tmp / "bad.xlsx"; shutil.copy(ROOT / "specs/cargo_shipments.xlsx", bad)
wb = openpyxl.load_workbook(bad); ws = wb["Columns"]
ws["A2"] = "AWB Number"; ws["E3"] = "maybe"; ws["I4"] = None; ws["C5"] = "text"
wb["Quality"]["E3"] = "CAST(substr(awb_number, 5, 7) AS BIGINT) > 0"
wb.save(bad)
rc, out = run(["convert", str(bad), str(tmp / "bad.yaml")])
assert rc == 1 and not (tmp / "bad.yaml").exists(), out
for needle in ("Columns!row 2", "awb_number", "Columns!row 3", "yes/no", "Columns!row 4", "description is required",
               "Columns!row 5", "logical_type 'text'"):
    assert needle in out, (needle, out)
# 6. template opens and has all sheets
rc, _ = run(["template", str(tmp / "t.xlsx")])
assert rc == 0 and {"README", "Contract", "Columns", "Quality", "Team", "Roles"} <= set(openpyxl.load_workbook(tmp / "t.xlsx").sheetnames)
print("all excel_contract tests passed")
