"""
excel_contract — turn a business-friendly Excel spec into an ODCS v3 data contract (team profile).

Commands (CLI or import):
  template      <out.xlsx>                                  blank spec workbook with instructions + dropdowns
  convert       <spec.xlsx> <contract.yaml> [--previous OLD.yaml] [--allow-breaking] [--dry-run]
                                                            Excel -> contract; validates; proposes/checks version
  from-contract <contract.yaml> <out.xlsx>                  existing contract -> Excel (round trip / onboarding)
  set           <spec.xlsx> <field> <value>                 change one field on the Contract sheet (e.g. catalog)

Exit code 0 = written (or nothing to do), 1 = blocking issues, 2 = usage error.

Versioning (when a previous contract exists — by default the output file itself):
  * the change is classified with contract-driven-pipeline's `diff` (METADATA / ADDITIVE / BREAKING)
  * version left blank in Excel -> proposed automatically (patch / minor / major)
  * version given but too small for the change -> blocked (semver violation)
  * BREAKING -> blocked unless --allow-breaking (the data owner has signed off)

In a Databricks notebook:
  %pip install -q openpyxl pyyaml
  import sys; sys.path.insert(0, "<skill folder>/scripts"); import excel_contract as xc
  xc.main(["convert", "<spec.xlsx>", "<contracts/name.odcs.yaml>"])
"""
from __future__ import annotations

import re
import sys
import uuid
from pathlib import Path

try:
    import yaml
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError as ex:  # pragma: no cover
    sys.exit(f"{ex.name} missing: run %pip install openpyxl pyyaml")

# contract_tools lives in the sibling skill; reuse its validator and diff so both skills share one profile
_HERE = Path(__file__).resolve()
for cand in (_HERE.parents[2] / "contract-driven-pipeline" / "scripts",):
    if (cand / "contract_tools.py").exists():
        sys.path.insert(0, str(cand))
try:
    import contract_tools as ct
except ImportError:  # pragma: no cover
    ct = None

SNAKE = re.compile(r"^[a-z][a-z0-9_]*$")
LOGICAL_DEFAULT_PHYSICAL = {"string": "STRING", "integer": "INT", "number": "DECIMAL(18,2)", "boolean": "BOOLEAN",
                            "date": "DATE", "timestamp": "TIMESTAMP"}
CLASSIFICATIONS = ["public", "internal", "confidential", "restricted"]
SEVERITIES = ["error", "warning", "critical"]
YES, NO = {"yes", "y", "true", "1", "x", "required", "ja"}, {"no", "n", "false", "0", "", "optional", "nein"}

# --------------------------------------------------------------------------- #
# Workbook layout (single source of truth for template, reader and writer)
# --------------------------------------------------------------------------- #
CONTRACT_FIELDS = [  # (field, required, help)
    ("name", True, "Dataset / contract name, snake_case (e.g. cargo_shipments). Becomes silver_<name>."),
    ("version", False, "Semantic version. Leave blank to let the tool propose it from the change."),
    ("id", False, "Stable contract id. Leave blank on first creation — the tool derives one."),
    ("status", False, "active | draft | proposed | deprecated | retired (default active)."),
    ("domain", True, "Business domain, e.g. cargo."),
    ("data_product", False, "Data product the dataset belongs to."),
    ("tenant", False, "Tenant / organisation."),
    ("purpose", True, "Why the dataset exists."),
    ("usage", False, "How consumers may use it (joins, restrictions)."),
    ("limitations", False, "Known limitations (late data, gaps)."),
    ("dataset_business_name", False, "Readable dataset name."),
    ("dataset_description", True, "One sentence: becomes the table comment."),
    ("dataset_granularity", False, "What one row represents."),
    ("dataset_tags", False, "Comma-separated, lowercase_with_underscores."),
    ("server", False, "Server label (default demo)."),
    ("environment", False, "dev | test | prod (default dev)."),
    ("catalog", True, "Unity Catalog catalog of the target tables."),
    ("schema", True, "Unity Catalog schema of the target tables."),
    ("source_format", False, "json | csv | parquet (default json)."),
    ("source_path_conf", True, "Pipeline configuration key that holds the landing path, e.g. cargo.landing_path."),
    ("source_path", True, "Landing location, e.g. /Volumes/<catalog>/<schema>/<volume>/<folder>."),
    ("latency_value", False, "Freshness target (number)."),
    ("latency_unit", False, "h | d | m (minutes)."),
    ("latency_column", False, "Column the freshness is measured on."),
    ("retention_value", False, "Retention (number)."),
    ("retention_unit", False, "d | m | y."),
]
COLUMN_HEADERS = ["name", "business_name", "logical_type", "physical_type", "required", "primary_key_position",
                  "classification", "tags", "description", "pattern", "minimum", "exclusive_minimum", "maximum",
                  "exclusive_maximum", "example"]
QUALITY_HEADERS = ["name", "description", "dimension", "severity", "sql_predicate"]
TEAM_HEADERS = ["username", "role"]
ROLE_HEADERS = ["role", "access"]


class Issue:
    def __init__(self, level, where, msg):
        self.level, self.where, self.msg = level, where, msg

    def __str__(self):
        return f"[{self.level}] {self.where}: {self.msg}"


# --------------------------------------------------------------------------- #
# reading helpers
# --------------------------------------------------------------------------- #
def _s(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return " ".join(str(v).split()) if not isinstance(v, str) else v.strip()


def _num(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    f = float(v)
    return int(f) if f.is_integer() else f


def _list(v) -> list[str]:
    return [t.strip() for t in re.split(r"[,;\n]", _s(v)) if t.strip()]


def _sheet(wb, name):
    for ws in wb.worksheets:
        if ws.title.strip().lower() == name.lower():
            return ws
    return None


def _table(ws, headers, issues, sheet):
    """Rows of a header table as dicts (header match is case/space-insensitive). Skips blank rows."""
    rows = list(ws.iter_rows(values_only=True))
    hdr_idx = None
    for i, r in enumerate(rows[:10]):
        cells = [_s(c).lower().replace(" ", "_") for c in r]
        if "name" in cells or "username" in cells or "role" in cells:
            hdr_idx, found = i, cells
            break
    if hdr_idx is None:
        issues.append(Issue("ERROR", sheet, f"header row not found (expected columns: {', '.join(headers)})"))
        return []
    pos = {h: found.index(h) for h in headers if h in found}
    missing = [h for h in headers if h not in pos]
    if missing:
        issues.append(Issue("WARNING", sheet, f"columns not found, treated as empty: {missing}"))
    out = []
    for n, r in enumerate(rows[hdr_idx + 1:], start=hdr_idx + 2):
        rec = {h: (r[pos[h]] if h in pos and pos[h] < len(r) else None) for h in headers}
        if all(_s(v) == "" for v in rec.values()):
            continue
        rec["_row"] = n
        out.append(rec)
    return out


def read_spec(path) -> tuple[dict, list[Issue]]:
    """Excel -> plain dict spec + issues found while reading."""
    issues: list[Issue] = []
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = _sheet(wb, "Contract")
    fields = {}
    if ws is None:
        issues.append(Issue("ERROR", "Contract", "sheet 'Contract' missing"))
    else:
        for r in ws.iter_rows(values_only=True):
            if r and _s(r[0]):
                fields[_s(r[0]).lower().replace(" ", "_")] = r[1] if len(r) > 1 else None
    spec = {"fields": fields, "columns": [], "quality": [], "team": [], "roles": []}
    for key, sheet, headers in (("columns", "Columns", COLUMN_HEADERS), ("quality", "Quality", QUALITY_HEADERS),
                                ("team", "Team", TEAM_HEADERS), ("roles", "Roles", ROLE_HEADERS)):
        ws = _sheet(wb, sheet)
        if ws is None:
            if key in ("columns",):
                issues.append(Issue("ERROR", sheet, f"sheet '{sheet}' missing"))
            continue
        spec[key] = _table(ws, headers, issues, sheet)
    return spec, issues


# --------------------------------------------------------------------------- #
# spec -> contract
# --------------------------------------------------------------------------- #
def build_contract(spec: dict, previous: dict | None = None) -> tuple[dict, list[Issue]]:
    f, issues = spec["fields"], []
    E = lambda w, m: issues.append(Issue("ERROR", w, m))
    W = lambda w, m: issues.append(Issue("WARNING", w, m))
    g = lambda k: _s(f.get(k))

    for k, req, _ in CONTRACT_FIELDS:
        if req and not g(k):
            E(f"Contract!{k}", "is required")
    name = g("name")
    if name and not SNAKE.match(name):
        E("Contract!name", f"'{name}' must be snake_case, e.g. '{re.sub(r'[^a-z0-9]+', '_', name.lower()).strip('_')}'")

    cid = g("id") or (previous or {}).get("id")
    if not cid and name:
        seed = "datacontract:" + g("domain") + ":" + name
        cid = uuid.uuid5(uuid.NAMESPACE_URL, seed).hex[:8] + "-" + name.replace("_", "-")
        W("Contract!id", f"derived '{cid}' — copy it into the Excel so it stays stable")

    c = {"apiVersion": "v3.0.2", "kind": "DataContract", "id": cid, "name": name,
         "version": g("version") or None, "status": g("status") or "active", "domain": g("domain")}
    for k, ck in (("data_product", "dataProduct"), ("tenant", "tenant")):
        if g(k):
            c[ck] = g(k)
    desc = {k: g(k) for k in ("purpose", "usage", "limitations") if g(k)}
    if desc:
        c["description"] = desc
    c["servers"] = [{"server": g("server") or "demo", "type": "databricks", "environment": g("environment") or "dev",
                     "catalog": g("catalog"), "schema": g("schema")}]

    team = [{"username": _s(t["username"]), "role": _s(t["role"])} for t in spec["team"] if _s(t["username"])]
    if team:
        c["team"] = team
    else:
        W("Team", "no owner listed")
    roles = [{"role": _s(r["role"]), "access": _s(r["access"]).lower()} for r in spec["roles"] if _s(r["role"])]
    if roles:
        c["roles"] = roles

    sla = []
    if g("latency_value"):
        item = {"property": "latency", "value": _num(f.get("latency_value")), "unit": g("latency_unit") or "h"}
        if g("latency_column"):
            item["element"] = f"{name}.{g('latency_column')}"
        sla.append(item)
    if g("retention_value"):
        sla.append({"property": "retention", "value": _num(f.get("retention_value")), "unit": g("retention_unit") or "y"})
    if sla:
        c["slaProperties"] = sla
    c["customProperties"] = [{"property": "sourceFormat", "value": g("source_format") or "json"},
                             {"property": "sourcePathConf", "value": g("source_path_conf")},
                             {"property": "sourcePath", "value": g("source_path")}]

    obj = {"name": name, "physicalName": f"silver_{name}", "physicalType": "table"}
    if g("dataset_business_name"):
        obj["businessName"] = g("dataset_business_name")
    obj["description"] = g("dataset_description")
    if g("dataset_granularity"):
        obj["dataGranularityDescription"] = g("dataset_granularity")
    if g("dataset_tags"):
        obj["tags"] = _list(f.get("dataset_tags"))

    props, seen = [], set()
    for r in spec["columns"]:
        w = f"Columns!row {r['_row']}"
        n = _s(r["name"])
        if not n:
            E(w, "name is empty"); continue
        if not SNAKE.match(n):
            E(w, f"column '{n}' must be snake_case, e.g. '{re.sub(r'[^a-z0-9]+', '_', n.lower()).strip('_')}'")
        if n in seen:
            E(w, f"duplicate column '{n}'")
        seen.add(n)
        lt = _s(r["logical_type"]).lower()
        if lt not in LOGICAL_DEFAULT_PHYSICAL:
            E(w, f"logical_type '{lt}' not in {list(LOGICAL_DEFAULT_PHYSICAL)}"); continue
        pt = _s(r["physical_type"]).upper().replace(" ", "") or LOGICAL_DEFAULT_PHYSICAL[lt]
        if not _s(r["physical_type"]):
            W(w, f"physical_type blank for '{n}', using {pt}")
        req_raw = _s(r["required"]).lower()
        if req_raw not in YES | NO:
            E(w, f"required must be yes/no, got '{r['required']}'")
        required = req_raw in YES
        cls = _s(r["classification"]).lower() or "internal"
        if not _s(r["classification"]):
            W(w, f"classification blank for '{n}', using internal")
        if cls not in CLASSIFICATIONS:
            E(w, f"classification '{cls}' not in {CLASSIFICATIONS}")
        p = {"name": n}
        if _s(r["business_name"]):
            p["businessName"] = _s(r["business_name"])
        p.update({"logicalType": lt, "physicalType": pt, "required": required})
        pk = _s(r["primary_key_position"])
        if pk:
            try:
                p["primaryKey"], p["primaryKeyPosition"] = True, int(float(pk))
            except ValueError:
                E(w, f"primary_key_position must be a number, got '{pk}'")
            if not required:
                E(w, f"primary key column '{n}' must be required")
        p["classification"] = cls
        tags = _list(r["tags"])
        if tags:
            p["tags"] = tags
        if not _s(r["description"]):
            E(w, f"description is required for '{n}' (it becomes the column comment)")
        p["description"] = _s(r["description"])
        opts = {}
        if _s(r["pattern"]):
            opts["pattern"] = _s(r["pattern"])
        for k, ok in (("minimum", "minimum"), ("exclusive_minimum", "exclusiveMinimum"),
                      ("maximum", "maximum"), ("exclusive_maximum", "exclusiveMaximum")):
            try:
                v = _num(r[k])
            except ValueError:
                E(w, f"{k} must be a number, got '{r[k]}'"); v = None
            if v is not None:
                opts[ok] = v
        if opts:
            p["logicalTypeOptions"] = opts
        if _s(r["example"]):
            p["examples"] = _list(r["example"])
        props.append(p)
    pks = sorted(p["primaryKeyPosition"] for p in props if p.get("primaryKey"))
    if pks and pks != list(range(1, len(pks) + 1)):
        E("Columns", f"primary_key_position values must be 1..n without gaps, got {pks}")
    obj["properties"] = props

    quality = []
    for r in spec["quality"]:
        w = f"Quality!row {r['_row']}"
        n = _s(r["name"])
        sev = _s(r["severity"]).lower() or "error"
        if sev not in SEVERITIES:
            E(w, f"severity '{sev}' not in {SEVERITIES}")
        if not _s(r["sql_predicate"]):
            E(w, f"rule '{n}' has no sql_predicate")
        q = {"name": n, "description": _s(r["description"])}
        if _s(r["dimension"]):
            q["dimension"] = _s(r["dimension"]).lower()
        q.update({"severity": sev, "type": "custom", "engine": "sql-predicate", "implementation": _s(r["sql_predicate"])})
        quality.append(q)
    if quality:
        obj["quality"] = quality
    c["schema"] = [obj]
    return c, issues


# --------------------------------------------------------------------------- #
# versioning against the previous contract
# --------------------------------------------------------------------------- #
def _bump(version: str, kind: str) -> str:
    ma, mi, pa = map(int, version.split("."))
    return {"major": f"{ma + 1}.0.0", "minor": f"{ma}.{mi + 1}.0", "patch": f"{ma}.{mi}.{pa + 1}"}[kind]


def resolve_version(c: dict, previous: dict | None, allow_breaking: bool) -> tuple[str, list[Issue], dict | None]:
    """Returns (status, issues, diff). status: create | update | unchanged | blocked."""
    issues = []
    if previous is None:
        c["version"] = c["version"] or "1.0.0"
        return "create", issues, None
    if ct is None:
        issues.append(Issue("ERROR", "versioning", "contract-driven-pipeline skill not found next to this skill"))
        return "blocked", issues, None
    probe = dict(c, version=c["version"] or str(previous["version"]))
    d = ct.diff_contracts(previous, probe)
    if d["overall"] == "NONE" and (not c["version"] or c["version"] == str(previous["version"])):
        c["version"] = str(previous["version"])
        return "unchanged", issues, d
    if not c["version"]:
        c["version"] = _bump(str(previous["version"]), d["required_bump"] if d["required_bump"] != "none" else "patch")
        issues.append(Issue("INFO", "version", f"proposed {c['version']} ({d['overall']} change, "
                                               f"{d['required_bump']} bump from {previous['version']})"))
        d = ct.diff_contracts(previous, c)
    else:
        d = ct.diff_contracts(previous, c)
        if not d["semver_ok"]:
            issues.append(Issue("ERROR", "version", f"{previous['version']} -> {c['version']} is a {d['actual_bump']} "
                                                    f"bump, but the change is {d['overall']} (needs {d['required_bump']})"))
            return "blocked", issues, d
    if d["overall"] == "BREAKING" and not allow_breaking:
        issues.append(Issue("ERROR", "breaking", "BREAKING change — not written. The data owner must confirm; "
                                                 "then re-run with --allow-breaking"))
        return "blocked", issues, d
    return "update", issues, d


# --------------------------------------------------------------------------- #
# writing
# --------------------------------------------------------------------------- #
class _Dumper(yaml.SafeDumper):
    pass


def _str_rep(dumper, data):
    style = ">" if len(data) > 100 and "\n" not in data else None
    return dumper.represent_scalar("tag:yaml.org,2002:str", data, style=style)


_Dumper.add_representer(str, _str_rep)


def to_yaml(c: dict, source: str) -> str:
    head = (f"# ODCS v3 data contract — GENERATED from {source} by the excel-to-data-contract skill.\n"
            f"# Edit the Excel and regenerate; hand edits here are overwritten.\n")
    return head + yaml.dump(c, Dumper=_Dumper, sort_keys=False, allow_unicode=True, width=110)


def convert(spec_path, out_path, previous_path=None, allow_breaking=False, dry_run=False) -> int:
    spec, issues = read_spec(spec_path)
    prev_p = Path(previous_path) if previous_path else Path(out_path)
    previous = yaml.safe_load(prev_p.read_text(encoding="utf-8")) if prev_p.exists() else None
    c, more = build_contract(spec, previous)
    issues += more
    if not any(i.level == "ERROR" for i in issues):
        status, vi, d = resolve_version(c, previous, allow_breaking)
        issues += vi
        if ct is not None and status != "blocked":
            issues += [Issue(x.level, x.where, x.message) for x in ct.validate_contract(c)]
    else:
        status, d = "blocked", None
    for i in issues:
        print(i)
    if d:
        print(f"Change vs {prev_p.name} v{d['from']}: {d['overall']}")
        for x in d["_changes"]:
            print("   ", x)
    errors = [i for i in issues if i.level == "ERROR"]
    if errors or status == "blocked":
        if any(i.where == "breaking" for i in errors):
            print("BLOCKED — breaking change. Nothing written. Needs the data owner's sign-off (then --allow-breaking).")
        else:
            print(f"BLOCKED — {len(errors)} error(s). Nothing written. Fix the Excel (cells are named above).")
        return 1
    if status == "unchanged":
        print(f"UNCHANGED — {Path(out_path).name} already matches the Excel (v{c['version']}). Nothing written.")
        return 0
    if dry_run:
        print(to_yaml(c, Path(spec_path).name))
        print(f"DRY RUN — would {status} {out_path} as v{c['version']}.")
        return 0
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    Path(out_path).write_text(to_yaml(c, Path(spec_path).name), encoding="utf-8")
    print(f"WRITTEN — {status} {out_path} v{c['version']} "
          f"({len(c['schema'][0]['properties'])} columns, {len(c['schema'][0].get('quality', []))} quality rules).")
    return 0


# --------------------------------------------------------------------------- #
# template / from-contract / set
# --------------------------------------------------------------------------- #
NAVY, LIGHT = "102C4D", "E1F0FB"


def _style_header(ws, ncols):
    for i in range(1, ncols + 1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"


def _workbook(contract: dict | None = None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "README"
    lines = [
        "Data contract specification — fill this workbook, then ask Genie Code:",
        "  @excel-to-data-contract Create the data contract from <this file>",
        "",
        "Sheets: Contract (one value per field) · Columns (one row per column) · Quality (business rules) · Team · Roles",
        "Required: grey-marked fields on Contract; every column needs name, logical_type, required, classification, description.",
        "Leave 'version' blank — the tool proposes it from the change (patch / minor / major).",
        "Patterns: prefer character classes ([0-9], [.]) over backslashes.",
        "Quality rules: one SQL boolean per row, true = row is valid. Use try_cast, never CAST.",
        "Severity: error = row rejected to quarantine · warning = kept, counted · critical = pipeline update stops.",
    ]
    for i, t in enumerate(lines, 1):
        ws.cell(row=i, column=1, value=t).font = Font(bold=(i == 1))
    ws.column_dimensions["A"].width = 120

    ws = wb.create_sheet("Contract")
    ws.append(["field", "value", "help"])
    _style_header(ws, 3)
    vals = _contract_to_fields(contract) if contract else {}
    for k, req, help_ in CONTRACT_FIELDS:
        ws.append([k, vals.get(k), ("REQUIRED · " if req else "") + help_])
        if req:
            ws.cell(row=ws.max_row, column=1).fill = PatternFill("solid", fgColor="D9D9D9")
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width, ws.column_dimensions["C"].width = 24, 60, 90
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")

    ws = wb.create_sheet("Columns")
    ws.append(COLUMN_HEADERS)
    _style_header(ws, len(COLUMN_HEADERS))
    widths = [22, 22, 13, 15, 10, 12, 14, 16, 60, 26, 10, 12, 10, 12, 18]
    for i, wdt in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = wdt
    for p in (contract["schema"][0]["properties"] if contract else []):
        o = p.get("logicalTypeOptions") or {}
        ws.append([p["name"], p.get("businessName"), p.get("logicalType"), p.get("physicalType"),
                   "yes" if p.get("required") else "no", p.get("primaryKeyPosition"), p.get("classification"),
                   ", ".join(p.get("tags") or []) or None, " ".join(str(p.get("description", "")).split()),
                   o.get("pattern"), o.get("minimum"), o.get("exclusiveMinimum"), o.get("maximum"),
                   o.get("exclusiveMaximum"), ", ".join(p.get("examples") or []) or None])
    for col, opts in (("C", list(LOGICAL_DEFAULT_PHYSICAL)), ("E", ["yes", "no"]), ("G", CLASSIFICATIONS)):
        dv = DataValidation(type="list", formula1='"' + ",".join(opts) + '"', allow_blank=True)
        ws.add_data_validation(dv); dv.add(f"{col}2:{col}500")

    ws = wb.create_sheet("Quality")
    ws.append(QUALITY_HEADERS)
    _style_header(ws, len(QUALITY_HEADERS))
    for col, wdt in zip("ABCDE", (30, 50, 14, 10, 90)):
        ws.column_dimensions[col].width = wdt
    for q in (contract["schema"][0].get("quality") or [] if contract else []):
        ws.append([q["name"], q.get("description"), q.get("dimension"), q.get("severity"),
                   " ".join(str(q.get("implementation", "")).split())])
    dv = DataValidation(type="list", formula1='"' + ",".join(SEVERITIES) + '"', allow_blank=True)
    ws.add_data_validation(dv); dv.add("D2:D500")

    for title, headers, rows in (("Team", TEAM_HEADERS, (contract or {}).get("team") or []),
                                 ("Roles", ROLE_HEADERS, (contract or {}).get("roles") or [])):
        ws = wb.create_sheet(title)
        ws.append(headers)
        _style_header(ws, len(headers))
        ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 40, 24
        for r in rows:
            ws.append([r.get(headers[0]), r.get(headers[1])])
    wb.active = 1
    return wb


def _contract_to_fields(c: dict) -> dict:
    o = c["schema"][0]
    srv = (c.get("servers") or [{}])[0]
    custom = {p["property"]: p["value"] for p in c.get("customProperties") or []}
    sla = {s["property"]: s for s in c.get("slaProperties") or []}
    d = c.get("description") or {}
    lat = sla.get("latency", {})
    return {
        "name": c.get("name"), "version": c.get("version"), "id": c.get("id"), "status": c.get("status"),
        "domain": c.get("domain"), "data_product": c.get("dataProduct"), "tenant": c.get("tenant"),
        "purpose": " ".join(str(d.get("purpose", "")).split()) or None,
        "usage": " ".join(str(d.get("usage", "")).split()) or None,
        "limitations": " ".join(str(d.get("limitations", "")).split()) or None,
        "dataset_business_name": o.get("businessName"), "dataset_description": o.get("description"),
        "dataset_granularity": o.get("dataGranularityDescription"), "dataset_tags": ", ".join(o.get("tags") or []) or None,
        "server": srv.get("server"), "environment": srv.get("environment"), "catalog": srv.get("catalog"),
        "schema": srv.get("schema"), "source_format": custom.get("sourceFormat"),
        "source_path_conf": custom.get("sourcePathConf"), "source_path": custom.get("sourcePath"),
        "latency_value": lat.get("value"), "latency_unit": lat.get("unit"),
        "latency_column": str(lat.get("element", "")).split(".")[-1] or None,
        "retention_value": sla.get("retention", {}).get("value"), "retention_unit": sla.get("retention", {}).get("unit"),
    }


def set_field(path, field, value):
    wb = openpyxl.load_workbook(path)
    ws = _sheet(wb, "Contract")
    for row in ws.iter_rows(min_row=1):
        if _s(row[0].value).lower() == field.lower():
            row[1].value = value
            wb.save(path)
            return True
    return False


def main(argv=None) -> int:
    argv = list(sys.argv[1:] if argv is None else argv)
    if not argv:
        print(__doc__); return 2
    cmd, args = argv[0], argv[1:]
    flag = lambda f: (f in args) and (args.remove(f) or True)

    def opt(name):
        if name in args:
            i = args.index(name); v = args[i + 1]; del args[i:i + 2]; return v
        return None

    if cmd == "template" and len(args) == 1:
        _workbook().save(args[0]); print(f"Wrote blank spec template {args[0]}"); return 0
    if cmd == "from-contract" and len(args) == 2:
        c = yaml.safe_load(Path(args[0]).read_text(encoding="utf-8"))
        _workbook(c).save(args[1]); print(f"Wrote {args[1]} from {args[0]} (v{c.get('version')})"); return 0
    if cmd == "set" and len(args) == 3:
        ok = set_field(args[0], args[1], args[2])
        print(("Set " if ok else "Field not found: ") + f"{args[1]} = {args[2]} in {args[0]}"); return 0 if ok else 1
    if cmd == "convert":
        prev, allow, dry = opt("--previous"), flag("--allow-breaking"), flag("--dry-run")
        if len(args) == 2:
            return convert(args[0], args[1], prev, allow, dry)
    print(__doc__)
    return 2


if __name__ == "__main__":
    sys.exit(main())
